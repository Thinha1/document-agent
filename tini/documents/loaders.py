"""Safe, local document parsing with optional PDF and DOCX adapters."""

from __future__ import annotations

import hashlib
import mimetypes
from dataclasses import dataclass, field
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from tini.documents.ocr import OcrProvider

PLAIN_TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".csv",
    ".json",
    ".yaml",
    ".yml",
}
SUPPORTED_SUFFIXES = PLAIN_TEXT_SUFFIXES | {".html", ".htm", ".pdf", ".docx"}


class DocumentLoadError(ValueError):
    """A source is unsupported, unsafe, too large, or cannot be parsed."""


@dataclass(frozen=True)
class LoadedDocument:
    path: Path
    source_uri: str
    title: str
    text: str
    media_type: str
    size_bytes: int
    content_hash: str
    metadata: dict[str, Any] = field(default_factory=dict)


class _TextOnlyHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"script", "style", "noscript"}:
            self._ignored_depth += 1
        elif tag in {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self._ignored_depth:
            self._ignored_depth -= 1
        elif tag in {"p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth:
            self.parts.append(data)

    def text(self) -> str:
        return "\n".join(line.strip() for line in "".join(self.parts).splitlines() if line.strip())


def discover_documents(path: Path, recursive: bool = True) -> list[Path]:
    """Return supported, regular files in deterministic order.

    Directory traversal never follows symlinks. An explicitly supplied symlink
    is rejected as well, so a watched ingestion folder cannot escape its root.
    """
    path = path.expanduser()
    if not path.exists():
        raise DocumentLoadError(f"source does not exist: {path}")
    if path.is_symlink():
        raise DocumentLoadError(f"symbolic links are not ingested: {path}")
    if path.is_file():
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise DocumentLoadError(f"unsupported document type: {path.suffix or '(none)'}")
        return [path.resolve()]
    if not path.is_dir():
        raise DocumentLoadError(f"source is not a regular file or directory: {path}")

    iterator = path.rglob("*") if recursive else path.glob("*")
    files = [
        item.resolve()
        for item in iterator
        if item.is_file()
        and not item.is_symlink()
        and not item.name.startswith("~$")
        and item.suffix.lower() in SUPPORTED_SUFFIXES
    ]
    return sorted(files, key=lambda item: str(item).lower())


def load_document(
    path: Path,
    max_file_bytes: int,
    max_text_chars: int = 5_000_000,
    ocr: OcrProvider | None = None,
) -> LoadedDocument:
    path = path.expanduser()
    if path.is_symlink():
        raise DocumentLoadError(f"symbolic links are not ingested: {path}")
    path = path.resolve()
    if not path.is_file():
        raise DocumentLoadError(f"source is not a regular file: {path}")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentLoadError(f"unsupported document type: {suffix or '(none)'}")
    size = path.stat().st_size
    if size > max_file_bytes:
        raise DocumentLoadError(
            f"document is {size} bytes; limit is {max_file_bytes} bytes: {path.name}"
        )

    with path.open("rb") as source_file:
        raw = source_file.read(max_file_bytes + 1)
    if len(raw) > max_file_bytes:
        raise DocumentLoadError(
            f"document exceeds the {max_file_bytes} byte limit while reading: {path.name}"
        )
    size = len(raw)
    content_hash = hashlib.sha256(raw).hexdigest()
    metadata: dict[str, Any] = {"extension": suffix}
    title = path.stem

    if suffix in PLAIN_TEXT_SUFFIXES:
        text = _decode_utf8(raw, path)
    elif suffix in {".html", ".htm"}:
        parser = _TextOnlyHTMLParser()
        parser.feed(_decode_utf8(raw, path))
        text = parser.text()
    elif suffix == ".pdf":
        text, pdf_meta = _load_pdf(raw, path, ocr=ocr)
        metadata.update(pdf_meta)
    else:
        text, docx_meta = _load_docx(raw, path)
        title = docx_meta.pop("title", "") or title
        metadata.update(docx_meta)

    text = _normalise_text(text)
    if not text:
        raise DocumentLoadError(f"document contains no extractable text: {path.name}")
    if len(text) > max_text_chars:
        raise DocumentLoadError(
            f"extracted text is {len(text)} characters; limit is {max_text_chars}: {path.name}"
        )
    media_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    return LoadedDocument(
        path=path,
        source_uri=path.as_uri(),
        title=title,
        text=text,
        media_type=media_type,
        size_bytes=size,
        content_hash=content_hash,
        metadata=metadata,
    )


def _decode_utf8(raw: bytes, path: Path) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentLoadError(f"text document must be UTF-8: {path.name}") from exc


def _normalise_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in text.splitlines()]
    return "\n".join(lines).strip()


# Below this average extracted characters per page, a PDF is treated as
# scanned/image-only rather than a normal text PDF. Calibrated against real
# documents: a genuinely image-only 15-page PDF averaged ~17 chars/page
# (stray glyphs pypdf still catches), while every text PDF sampled averaged
# in the hundreds. 40 sits well inside that gap.
_MIN_CHARS_PER_PAGE = 40


def _load_pdf(
    raw: bytes, path: Path, ocr: OcrProvider | None = None
) -> tuple[str, dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentLoadError(
            "PDF ingestion requires pypdf; reinstall tini-agent to restore core dependencies"
        ) from exc
    try:
        reader = PdfReader(BytesIO(raw))
        # Many official/government PDFs (e.g. Japan's EDINET securities
        # filings) are "encrypted" only to restrict printing/copying in a
        # viewer, with an empty user password — reading is not actually
        # gated. pypdf's is_encrypted stays True even after a successful
        # decrypt (it reflects the file format, not decryption state), so
        # the empty-password attempt's own return value is the real check.
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentLoadError(f"encrypted PDF is not supported: {path.name}")
        pages = [page.extract_text() or "" for page in reader.pages]
    except DocumentLoadError:
        raise
    except Exception as exc:
        raise DocumentLoadError(f"could not parse PDF {path.name}: {exc}") from exc

    page_count = len(pages)
    text = "\n\n".join(pages)
    metadata: dict[str, Any] = {"page_count": page_count}
    scanned = page_count > 0 and len(text) / page_count < _MIN_CHARS_PER_PAGE
    if scanned:
        if ocr is None:
            # Honest, not silent: an operator inspecting this document's
            # metadata (dashboard, `tini documents`) can see WHY a document
            # that "ingested fine" has almost no content, rather than the
            # near-empty chunk just looking like a thin source document.
            metadata["ocr_skipped"] = (
                "page text looks scanned/image-only; no OCR provider configured "
                "(TINI_OCR_PROVIDER, see tini/documents/ocr.py)"
            )
        else:
            before = _usage_of(ocr)
            ocr_text, failed_pages = _ocr_pdf_pages(raw, ocr)
            after = _usage_of(ocr)
            if before is not None and after is not None:
                # What THIS document cost, in tokens (never a price — see
                # tini/ops/ocr_cost.py). Recorded even when the run yields no
                # text: a failed scan still spent the calls.
                metadata["ocr_usage"] = {
                    "provider": getattr(ocr, "provider_name", ""),
                    "model": getattr(ocr, "model", ""),
                    "pages": after["calls"] - before["calls"],
                    "cached_pages": after.get("cache_hits", 0) - before.get("cache_hits", 0),
                    "prompt_tokens": after["prompt_tokens"] - before["prompt_tokens"],
                    "completion_tokens": after["completion_tokens"] - before["completion_tokens"],
                }
            if ocr_text.strip():
                text = ocr_text
                metadata["ocr_applied"] = True
                if failed_pages:
                    metadata["ocr_failed_pages"] = failed_pages
    return text, metadata


def _usage_of(ocr: OcrProvider) -> dict[str, int] | None:
    snapshot = getattr(ocr, "usage_snapshot", None)
    return snapshot() if callable(snapshot) else None


def _ocr_pdf_pages(raw: bytes, ocr: OcrProvider) -> tuple[str, int]:
    """Render each page to a bitmap — exactly what a viewer would show — and
    run `ocr` over it, stitching results back together in page order.

    Tried pulling the page's embedded image out via pypdf first
    (page.images), since a scanned page is virtually always stored as one
    full-page raster XObject and that would have needed no rasterizer at
    all. Abandoned it against this corpus's real scanned PDFs: pypdf reported
    "Image and mask size not matching" and returned a broken 2x2-pixel image
    for the genuine scan on most pages (the real content lives in a
    differently-sized soft mask pypdf doesn't reconcile), silently losing
    the page's content rather than erroring — the exact failure mode a
    pipeline that must work "without manual editing" cannot have. Rendering
    the fully composed page with pypdfium2 (PDFium, the engine Chrome uses)
    sidesteps the mask-decoding question entirely: whatever a human would
    see on the page is what gets rasterized and sent to the model.

    Rendering stays single-threaded — libpdfium is not thread-safe for
    concurrent page access — but the OCR call is a multi-second network
    round trip and pages are independent, so THAT part runs up to
    `ocr.concurrency` pages at once (`ThreadPoolExecutor.map` returns
    results in submission order regardless of completion order, so page
    order survives the reordering that concurrency introduces).

    Returns (text, failed_page_count): one page's OCR call failing does not
    sink the whole document, but the count stays visible in the caller's
    metadata rather than vanishing silently."""
    from concurrent.futures import ThreadPoolExecutor
    from io import BytesIO as _BytesIO

    import pypdfium2 as pdfium

    from tini.documents.ocr import OcrError

    pdf = pdfium.PdfDocument(raw)
    try:
        images: list[bytes] = []
        for page in pdf:
            try:
                # scale=2.0 ≈ 144 DPI — legible for a vision model without
                # the request bodies of a full 300 DPI scan.
                image = page.render(scale=2.0).to_pil()
                buf = _BytesIO()
                image.convert("RGB").save(buf, format="PNG")
                images.append(buf.getvalue())
            finally:
                page.close()
    finally:
        pdf.close()

    def _transcribe(image_bytes: bytes) -> str | None:
        try:
            return ocr.ocr_image(image_bytes, "image/png")
        except OcrError:
            return None

    concurrency = max(1, min(int(getattr(ocr, "concurrency", 1)), len(images) or 1))
    if concurrency <= 1:
        results = [_transcribe(image) for image in images]
    else:
        with ThreadPoolExecutor(max_workers=concurrency) as pool:
            results = list(pool.map(_transcribe, images))

    failed = sum(1 for result in results if result is None)
    parts = [text for text in results if text]
    return "\n\n".join(parts), failed


def _load_docx(raw: bytes, path: Path) -> tuple[str, dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentLoadError(
            "DOCX ingestion requires python-docx; reinstall tini-agent to restore core dependencies"
        ) from exc
    try:
        document = Document(BytesIO(raw))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        title = document.core_properties.title or ""
    except Exception as exc:
        raise DocumentLoadError(f"could not parse DOCX {path.name}: {exc}") from exc
    return "\n\n".join(parts), {"title": title}
